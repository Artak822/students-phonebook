import io
import re
from datetime import datetime, timezone

from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import (
    ActiveIllnessNotFound,
    AlreadySick,
    EmployeeNotFound,
    GroupNotFound,
    IllnessNotFound,
    PhoneAlreadyExists,
    RoomFull,
    RoomNotFound,
)
from app.employees import history as hist
from app.admins.models import Admin
from app.employees.models import Employee, EmployeeChangeHistory, EmployeeIllness, EmployeeStatement, IllnessNote
from app.employees.schemas import EmployeeCreate, EmployeeUpdate, HistoryEntry, IllnessCreate, IllnessNoteCreate, IllnessRecover, IllnessUpdate, StatementCreate
from app.groups.models import Group
from app.rooms.models import Room


async def _get_or_404(db: AsyncSession, emp_id: int) -> Employee:
    result = await db.execute(select(Employee).where(Employee.id == emp_id, Employee.deleted_at.is_(None)))
    emp = result.scalar_one_or_none()
    if not emp:
        raise EmployeeNotFound()
    return emp


async def list_employees(
    db: AsyncSession,
    search: str | None,
    group_id: int | None,
    building: int | None,
    room_id: int | None,
    sick: bool | None,
    page: int,
    limit: int,
) -> dict:
    if limit > 200:
        from fastapi import HTTPException

        raise HTTPException(status_code=422, detail=[{"field": "limit", "message": "limit не может превышать 200"}])

    q = select(Employee).where(Employee.deleted_at.is_(None))

    if search:
        phone_digits = re.sub(r"\D", "", search)
        if phone_digits and len(phone_digits) >= 4:
            # search by phone digits OR fio
            normalized = phone_digits
            if len(phone_digits) == 11 and phone_digits[0] in ("7", "8"):
                normalized = "7" + phone_digits[1:]
            q = q.where(
                Employee.fio.ilike(f"%{search}%")
                | Employee.phone.ilike(f"%{normalized}%")
            )
        else:
            q = q.where(Employee.fio.ilike(f"%{search}%"))

    if group_id is not None:
        q = q.where(Employee.group_id == group_id)
    if building is not None:
        room_ids_q = select(Room.id).where(Room.building == building)
        room_ids_result = await db.execute(room_ids_q)
        room_ids = [r for (r,) in room_ids_result.all()]
        q = q.where(Employee.room_id.in_(room_ids))
    if room_id is not None:
        q = q.where(Employee.room_id == room_id)
    if sick is True:
        sick_ids_q = select(EmployeeIllness.employee_id).where(EmployeeIllness.end_date.is_(None))
        q = q.where(Employee.id.in_(sick_ids_q))

    count_result = await db.execute(select(func.count()).select_from(q.subquery()))
    total = count_result.scalar()

    offset = (page - 1) * limit
    q = q.order_by(Employee.fio).offset(offset).limit(limit)
    result = await db.execute(q)
    items = list(result.scalars().all())

    total_pages = max(1, (total + limit - 1) // limit)
    return {"items": items, "total": total, "page": page, "limit": limit, "total_pages": total_pages}


async def get_employee(db: AsyncSession, emp_id: int) -> Employee:
    return await _get_or_404(db, emp_id)


async def create_employee(db: AsyncSession, data: EmployeeCreate, admin_id: int) -> Employee:
    # Check group exists
    group = await db.execute(select(Group).where(Group.id == data.group_id))
    if not group.scalar_one_or_none():
        raise GroupNotFound()
    # Check phone uniqueness
    existing = await db.execute(
        select(Employee).where(Employee.phone == data.phone, Employee.deleted_at.is_(None))
    )
    if existing.scalar_one_or_none():
        raise PhoneAlreadyExists()

    emp = Employee(
        fio=data.fio,
        phone=data.phone,
        group_id=data.group_id,
        birth_date=data.birth_date,
        notes=data.notes,
    )
    db.add(emp)
    await db.flush()
    await hist.record_create(db, emp.id, admin_id)
    await db.refresh(emp)
    return emp


async def update_employee(db: AsyncSession, emp_id: int, data: EmployeeUpdate, admin_id: int) -> Employee:
    emp = await _get_or_404(db, emp_id)

    update_data = data.model_dump(exclude_none=True)
    if not update_data:
        return emp

    if "phone" in update_data and update_data["phone"] != emp.phone:
        existing = await db.execute(
            select(Employee).where(
                Employee.phone == update_data["phone"],
                Employee.deleted_at.is_(None),
                Employee.id != emp_id,
            )
        )
        if existing.scalar_one_or_none():
            raise PhoneAlreadyExists()

    if "group_id" in update_data:
        group = await db.execute(select(Group).where(Group.id == update_data["group_id"]))
        if not group.scalar_one_or_none():
            raise GroupNotFound()

    old_values = {f: getattr(emp, f) for f in update_data}

    for field, value in update_data.items():
        setattr(emp, field, value)
    emp.updated_at = datetime.now(timezone.utc)

    await db.flush()
    await hist.record_update(db, emp.id, admin_id, old_values, update_data)
    await db.refresh(emp)
    return emp


async def soft_delete(db: AsyncSession, emp_id: int, admin_id: int) -> None:
    emp = await _get_or_404(db, emp_id)
    emp.deleted_at = datetime.now(timezone.utc)
    await db.flush()
    await hist.record_delete(db, emp.id, admin_id)


async def assign_room(db: AsyncSession, emp_id: int, room_id: int | None, admin_id: int) -> Employee:
    emp = await _get_or_404(db, emp_id)

    if room_id is not None:
        room_row = await db.execute(select(Room).where(Room.id == room_id))
        room = room_row.scalar_one_or_none()
        if not room:
            raise RoomNotFound()
        # Проверяем вместимость (не считаем самого студента, если он уже в этой комнате)
        occupants_q = select(func.count()).where(
            Employee.room_id == room_id,
            Employee.deleted_at.is_(None),
            Employee.id != emp_id,
        )
        occupants = (await db.execute(occupants_q)).scalar_one()
        if occupants >= room.capacity:
            raise RoomFull()

    old_room = emp.room_id
    emp.room_id = room_id
    emp.updated_at = datetime.now(timezone.utc)
    await db.flush()
    await hist.record_update(db, emp.id, admin_id, {"room_id": old_room}, {"room_id": room_id})
    await db.refresh(emp)
    return emp


async def get_roommates(db: AsyncSession, emp_id: int) -> list[Employee]:
    emp = await _get_or_404(db, emp_id)
    if emp.room_id is None:
        return []
    result = await db.execute(
        select(Employee).where(
            Employee.room_id == emp.room_id,
            Employee.id != emp_id,
            Employee.deleted_at.is_(None),
        )
    )
    return list(result.scalars().all())


async def list_statements(db: AsyncSession, emp_id: int) -> list[EmployeeStatement]:
    await _get_or_404(db, emp_id)
    result = await db.execute(
        select(EmployeeStatement)
        .where(EmployeeStatement.employee_id == emp_id)
        .order_by(EmployeeStatement.start_date.desc())
    )
    return list(result.scalars().all())


async def create_statement(db: AsyncSession, emp_id: int, data: StatementCreate) -> EmployeeStatement:
    await _get_or_404(db, emp_id)
    stmt = EmployeeStatement(
        employee_id=emp_id,
        start_date=data.start_date,
        end_date=data.end_date,
    )
    db.add(stmt)
    await db.flush()
    await db.refresh(stmt)
    return stmt


async def delete_statement(db: AsyncSession, emp_id: int, stmt_id: int) -> None:
    await _get_or_404(db, emp_id)
    result = await db.execute(
        select(EmployeeStatement).where(
            EmployeeStatement.id == stmt_id,
            EmployeeStatement.employee_id == emp_id,
        )
    )
    stmt = result.scalar_one_or_none()
    if stmt:
        await db.delete(stmt)
        await db.flush()


async def get_history(db: AsyncSession, emp_id: int) -> list[HistoryEntry]:
    await _get_or_404(db, emp_id)

    result = await db.execute(
        select(EmployeeChangeHistory, Admin.fio.label("admin_fio"))
        .outerjoin(Admin, Admin.id == EmployeeChangeHistory.admin_id)
        .where(EmployeeChangeHistory.employee_id == emp_id)
        .order_by(EmployeeChangeHistory.changed_at.desc())
    )
    rows = result.all()

    # Collect room IDs that need to be resolved
    room_ids: set[int] = set()
    for row in rows:
        h = row.EmployeeChangeHistory
        if h.field_name == "room_id":
            for val in (h.old_value, h.new_value):
                if val and val.isdigit():
                    room_ids.add(int(val))

    room_labels: dict[int, str] = {}
    if room_ids:
        room_result = await db.execute(
            select(Room.id, Room.building, Room.entrance, Room.room_number)
            .where(Room.id.in_(room_ids))
        )
        for r in room_result.all():
            room_labels[r.id] = f"{r.building}-{r.entrance}-{r.room_number}"

    def resolve_room(val: str | None) -> str | None:
        if not val:
            return None
        if val.isdigit():
            return room_labels.get(int(val), f"комната #{val}")
        return val

    return [
        HistoryEntry(
            id=row.EmployeeChangeHistory.id,
            employee_id=row.EmployeeChangeHistory.employee_id,
            admin_id=row.EmployeeChangeHistory.admin_id,
            admin_fio=row.admin_fio,
            action=row.EmployeeChangeHistory.action,
            field_name=row.EmployeeChangeHistory.field_name,
            old_value=resolve_room(row.EmployeeChangeHistory.old_value)
                if row.EmployeeChangeHistory.field_name == "room_id"
                else row.EmployeeChangeHistory.old_value,
            new_value=resolve_room(row.EmployeeChangeHistory.new_value)
                if row.EmployeeChangeHistory.field_name == "room_id"
                else row.EmployeeChangeHistory.new_value,
            changed_at=row.EmployeeChangeHistory.changed_at,
        )
        for row in rows
    ]


# ── Illnesses ─────────────────────────────────────────────────────────────────

async def _load_illness_with_notes(db: AsyncSession, illness: EmployeeIllness) -> EmployeeIllness:
    notes_result = await db.execute(
        select(IllnessNote)
        .where(IllnessNote.illness_id == illness.id)
        .order_by(IllnessNote.date.asc(), IllnessNote.created_at.asc())
    )
    illness.notes = list(notes_result.scalars().all())  # type: ignore[attr-defined]
    return illness


async def list_illnesses(db: AsyncSession, emp_id: int) -> list[EmployeeIllness]:
    await _get_or_404(db, emp_id)
    result = await db.execute(
        select(EmployeeIllness)
        .where(EmployeeIllness.employee_id == emp_id)
        .order_by(EmployeeIllness.start_date.desc())
    )
    illnesses = list(result.scalars().all())
    for ill in illnesses:
        await _load_illness_with_notes(db, ill)
    return illnesses


async def get_active_illness(db: AsyncSession, emp_id: int) -> EmployeeIllness | None:
    result = await db.execute(
        select(EmployeeIllness)
        .where(EmployeeIllness.employee_id == emp_id, EmployeeIllness.end_date.is_(None))
        .order_by(EmployeeIllness.start_date.desc())
        .limit(1)
    )
    ill = result.scalar_one_or_none()
    if ill:
        await _load_illness_with_notes(db, ill)
    return ill


async def update_illness(db: AsyncSession, emp_id: int, illness_id: int, data: IllnessUpdate) -> EmployeeIllness:
    await _get_or_404(db, emp_id)
    result = await db.execute(
        select(EmployeeIllness).where(EmployeeIllness.id == illness_id, EmployeeIllness.employee_id == emp_id)
    )
    illness = result.scalar_one_or_none()
    if not illness:
        raise IllnessNotFound()
    illness.temp_room_id = data.temp_room_id
    await db.flush()
    await _load_illness_with_notes(db, illness)
    return illness


async def create_illness(db: AsyncSession, emp_id: int, data: IllnessCreate, admin_id: int) -> EmployeeIllness:
    await _get_or_404(db, emp_id)
    # Can't have two active illnesses
    existing = await db.execute(
        select(EmployeeIllness)
        .where(EmployeeIllness.employee_id == emp_id, EmployeeIllness.end_date.is_(None))
    )
    if existing.scalar_one_or_none():
        raise AlreadySick()

    illness = EmployeeIllness(
        employee_id=emp_id,
        temp_room_id=data.temp_room_id,
        start_date=data.start_date,
    )
    db.add(illness)
    await db.flush()

    if data.first_note.strip():
        note = IllnessNote(
            illness_id=illness.id,
            admin_id=admin_id,
            note=data.first_note.strip(),
            date=data.start_date,
        )
        db.add(note)
        await db.flush()

    await _load_illness_with_notes(db, illness)
    return illness


async def add_illness_note(
    db: AsyncSession, emp_id: int, illness_id: int, data: IllnessNoteCreate, admin_id: int
) -> IllnessNote:
    await _get_or_404(db, emp_id)
    result = await db.execute(
        select(EmployeeIllness)
        .where(EmployeeIllness.id == illness_id, EmployeeIllness.employee_id == emp_id)
    )
    illness = result.scalar_one_or_none()
    if not illness:
        raise IllnessNotFound()

    note = IllnessNote(
        illness_id=illness_id,
        admin_id=admin_id,
        note=data.note.strip(),
        date=data.date,
    )
    db.add(note)
    await db.flush()
    await db.refresh(note)
    return note


async def delete_illness(db: AsyncSession, emp_id: int, illness_id: int) -> None:
    await _get_or_404(db, emp_id)
    result = await db.execute(
        select(EmployeeIllness).where(EmployeeIllness.id == illness_id, EmployeeIllness.employee_id == emp_id)
    )
    illness = result.scalar_one_or_none()
    if not illness:
        raise IllnessNotFound()
    # Delete notes first
    notes = await db.execute(select(IllnessNote).where(IllnessNote.illness_id == illness_id))
    for note in notes.scalars().all():
        await db.delete(note)
    await db.delete(illness)
    await db.flush()


async def export_employees_xlsx(
    db: AsyncSession,
    search: str | None,
    group_id: int | None,
    building: int | None,
    entrance: int | None,
    room_id: int | None,
    sick: bool | None,
) -> io.BytesIO:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    q = select(Employee).where(Employee.deleted_at.is_(None))

    if search:
        phone_digits = re.sub(r"\D", "", search)
        if phone_digits and len(phone_digits) >= 4:
            normalized = phone_digits
            if len(phone_digits) == 11 and phone_digits[0] in ("7", "8"):
                normalized = "7" + phone_digits[1:]
            q = q.where(
                Employee.fio.ilike(f"%{search}%") | Employee.phone.ilike(f"%{normalized}%")
            )
        else:
            q = q.where(Employee.fio.ilike(f"%{search}%"))

    if group_id is not None:
        q = q.where(Employee.group_id == group_id)
    if building is not None:
        room_ids_q = select(Room.id).where(Room.building == building)
        if entrance is not None:
            room_ids_q = room_ids_q.where(Room.entrance == entrance)
        room_ids_result = await db.execute(room_ids_q)
        room_ids_list = [r for (r,) in room_ids_result.all()]
        q = q.where(Employee.room_id.in_(room_ids_list))
    elif entrance is not None:
        room_ids_q = select(Room.id).where(Room.entrance == entrance)
        room_ids_result = await db.execute(room_ids_q)
        room_ids_list = [r for (r,) in room_ids_result.all()]
        q = q.where(Employee.room_id.in_(room_ids_list))
    if room_id is not None:
        q = q.where(Employee.room_id == room_id)
    if sick is True:
        sick_ids_q = select(EmployeeIllness.employee_id).where(EmployeeIllness.end_date.is_(None))
        q = q.where(Employee.id.in_(sick_ids_q))

    q = q.order_by(Employee.fio)
    result = await db.execute(q)
    employees = list(result.scalars().all())

    # Resolve group names
    group_ids = {e.group_id for e in employees}
    group_map: dict[int, str] = {}
    if group_ids:
        from app.groups.models import Group
        gr = await db.execute(select(Group.id, Group.name).where(Group.id.in_(group_ids)))
        group_map = {r.id: r.name for r in gr.all()}

    # Resolve room labels
    r_ids = {e.room_id for e in employees if e.room_id}
    room_map: dict[int, str] = {}
    if r_ids:
        rr = await db.execute(
            select(Room.id, Room.building, Room.entrance, Room.room_number).where(Room.id.in_(r_ids))
        )
        room_map = {r.id: f"{r.building}-{r.entrance}-{r.room_number}" for r in rr.all()}

    # Resolve sick status
    sick_ids: set[int] = set()
    if employees:
        emp_ids = [e.id for e in employees]
        sick_q = select(EmployeeIllness.employee_id).where(
            EmployeeIllness.employee_id.in_(emp_ids),
            EmployeeIllness.end_date.is_(None),
        )
        sk = await db.execute(sick_q)
        sick_ids = {r for (r,) in sk.all()}

    # Build workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Студенты"

    HEADERS = ["№", "ФИО", "Группа", "Дата рождения", "Комната", "Телефон", "Заметки", "На больничном"]
    COL_WIDTHS = [5, 36, 18, 16, 14, 18, 30, 14]

    header_fill = PatternFill("solid", fgColor="2C4A6E")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    thin = Side(style="thin", color="D0D8E4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    wrap = Alignment(wrap_text=True, vertical="top")

    for col_idx, (header, width) in enumerate(zip(HEADERS, COL_WIDTHS), 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = center
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 22

    alt_fill = PatternFill("solid", fgColor="F4F6FA")

    for row_idx, emp in enumerate(employees, 2):
        is_alt = row_idx % 2 == 0
        row_fill = PatternFill("solid", fgColor="F4F6FA") if is_alt else None
        bd = emp.birth_date
        bd_str = f"{bd.day:02d}.{bd.month:02d}.{bd.year}" if bd else ""

        values = [
            row_idx - 1,
            emp.fio,
            group_map.get(emp.group_id, ""),
            bd_str,
            room_map.get(emp.room_id, "") if emp.room_id else "",
            emp.phone,
            emp.notes or "",
            "Да" if emp.id in sick_ids else "Нет",
        ]

        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = border
            cell.alignment = wrap if col_idx == 7 else (center if col_idx in (1, 8) else Alignment(vertical="top"))
            if row_fill:
                cell.fill = row_fill

        ws.row_dimensions[row_idx].height = 18

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


async def recover_illness(
    db: AsyncSession, emp_id: int, illness_id: int, data: IllnessRecover, admin_id: int
) -> EmployeeIllness:
    await _get_or_404(db, emp_id)
    result = await db.execute(
        select(EmployeeIllness)
        .where(
            EmployeeIllness.id == illness_id,
            EmployeeIllness.employee_id == emp_id,
            EmployeeIllness.end_date.is_(None),
        )
    )
    illness = result.scalar_one_or_none()
    if not illness:
        raise ActiveIllnessNotFound()

    illness.end_date = data.end_date

    if data.final_note.strip():
        note = IllnessNote(
            illness_id=illness_id,
            admin_id=admin_id,
            note=data.final_note.strip(),
            date=data.end_date,
        )
        db.add(note)
        await db.flush()

    await _load_illness_with_notes(db, illness)
    return illness
